import io

from openpyxl import load_workbook

import exporters


def test_excel_splits_large_results_without_losing_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(exporters, "_EXCEL_MAX_ROWS", 4)
    data = exporters.reviews_to_bytes([
        {"platform": "google_play", "app_id": "app", "review_id": str(i), "text": f"Review {i}"}
        for i in range(7)
    ], "xlsx")
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    assert workbook.sheetnames == ["reviews", "reviews_2", "reviews_3"]
    values = []
    for sheet in workbook:
        rows = list(sheet.values)
        index = rows[0].index("review_id")
        values.extend(row[index] for row in rows[1:])
    assert values == [str(i) for i in range(7)]
    workbook.close()
